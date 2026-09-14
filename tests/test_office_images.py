"""
Office 文档内嵌图片提取测试

验证 convert_to_markdown 对 Office 文档的图片处理：
- 内嵌图片落盘到工作文件同名目录，引用改写为相对路径
- 同一张图按内容去重
- 正文中不残留 base64，也不残留 MarkItDown 的悬空占位名
- 不给落盘目录时正文同样不含 base64
"""

import sys
from pathlib import Path

import pytest
from PIL import Image as PILImage

# 添加项目根目录到 Python 路径
sys.path.insert(0, str(Path(__file__).parent.parent))

from indexing.services.converter import convert_to_markdown, extract_embedded_images

pptx = pytest.importorskip("pptx", reason="python-pptx 随 markitdown[all] 安装")


def _make_pptx(tmp_path: Path) -> Path:
    """构造 3 页 PPT，前两页复用同一张图用于验证去重。"""
    from pptx.util import Inches

    first = tmp_path / "first.png"
    PILImage.new("RGB", (400, 240), (30, 90, 180)).save(first)
    second = tmp_path / "second.png"
    PILImage.new("RGB", (300, 300), (200, 40, 40)).save(second)

    prs = pptx.Presentation()
    for index, picture in enumerate([first, first, second], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"第{index}页"
        slide.shapes.add_picture(str(picture), Inches(1), Inches(2), Inches(3))

    path = tmp_path / "演示文稿.pptx"
    prs.save(path)
    return path


def test_images_saved_and_deduped(tmp_path):
    """内嵌图片落盘到同名目录，相同内容只存一份"""
    source = _make_pptx(tmp_path)
    image_dir = tmp_path / "working" / "演示文稿"

    content = convert_to_markdown(source, image_dir)

    # 两张不同的图，前两页复用的那张不重复落盘
    assert sorted(p.suffix for p in image_dir.iterdir()) == [".png", ".png"]
    # 引用相对工作目录，形态与 OCR 解析产出一致
    assert content.count("(演示文稿/image-") == 3


def test_no_base64_or_dangling_placeholder(tmp_path):
    """正文既不残留 base64，也不残留 MarkItDown 的悬空占位名"""
    source = _make_pptx(tmp_path)

    content = convert_to_markdown(source, tmp_path / "working" / "演示文稿")

    assert "base64" not in content
    # keep_data_uris 关闭时 MarkItDown 会产出 Picture2.jpg 这类不存在的文件名
    assert ".jpg)" not in content


def test_without_image_dir_keeps_text_clean(tmp_path):
    """不提取图片时，正文也不应混入 base64"""
    source = _make_pptx(tmp_path)

    content = convert_to_markdown(source)

    assert "base64" not in content


def test_undecodable_data_uri_is_stripped(tmp_path):
    """无法解码的 data URI 必须剥离，避免 base64 进入切片"""
    image_dir = tmp_path / "working" / "文档"

    content = extract_embedded_images(
        "![图](data:image/png;base64,!!!非法!!!)", image_dir
    )

    assert "base64" not in content
    assert not image_dir.exists()
